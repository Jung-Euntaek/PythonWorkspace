from openpyxl import load_workbook
wb = load_workbook("sample.xlsx")
ws = wb.active

from openpyxl.chart import BarChart, Reference, LineChart
# B2:C11 데이터 차트로 생성
# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
# bar_chart = BarChart()
# bar_chart.add_data(bar_value)
# ws.add_chart(bar_chart, "E1")

line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data=True) # titles_from_data : 범례를 제목에서 가져옴
line_chart.title = "성적표"
line_chart.style = 20
line_chart.y_axis.title = "점수"
line_chart.x_axis.title = "번호"
ws.add_chart(line_chart, "E1")


wb.save("sample_chart.xlsx")